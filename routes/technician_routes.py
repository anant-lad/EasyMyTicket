"""
Technician Assistance Routes + Availability / Attendance
"""
import io
import base64
import logging
from datetime import datetime, timezone
from fastapi import APIRouter, File, Form, HTTPException, Depends, UploadFile
from pydantic import BaseModel, Field
from typing import Dict, Any, List, Optional
from src.database.db_connection import DatabaseConnection
from src.agents.technician_assistant import TechnicianAssistantAgent
from src.auth.dependencies import require_tech

log = logging.getLogger(__name__)

router = APIRouter()

# Lazy loading
_db_conn = None
_assistant_agent = None

def get_db_connection():
    global _db_conn
    if _db_conn is None:
        _db_conn = DatabaseConnection()
    return _db_conn

def get_assistant_agent():
    global _assistant_agent
    if _assistant_agent is None:
        _assistant_agent = TechnicianAssistantAgent(get_db_connection())
    return _assistant_agent

# Pydantic models
class TechnicianAssistRequest(BaseModel):
    text: str = Field(..., description="Natural language input from technician")
    session_id: Optional[str] = Field(None, description="Optional session ID for conversational context")

    model_config = {
        "json_schema_extra": {
            "example": {
                "text": "Help me with ticket T20240108.123456. I'm seeing a database connection error.",
                "session_id": "550e8400-e29b-41d4-a716-446655440000",
            }
        }
    }

class Source(BaseModel):
    ticket_number: str
    reason: str

class TechnicianAssistResponse(BaseModel):
    success: bool
    session_id: Optional[str] = None
    ticket_number: Optional[str] = None
    analysis: Optional[str] = None
    solution: Optional[str] = None
    sources: List[Source] = []
    follow_up_questions: List[str] = []
    message: Optional[str] = None
    original_query: Optional[str] = None

def _extract_file_text(filename: str, data: bytes, mime_type: str) -> str:
    """Extract plain text from an uploaded file for LLM context."""
    try:
        if mime_type == "text/plain" or mime_type == "text/csv":
            return data.decode("utf-8", errors="replace")

        if mime_type == "application/pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(data))
                return "\n".join(page.extract_text() or "" for page in reader.pages)
            except ImportError:
                return f"[PDF: {filename} — pypdf not installed]"

        ext = (filename or "").rsplit(".", 1)[-1].lower()

        if ext in ("docx",) or "wordprocessingml" in mime_type:
            try:
                from docx import Document
                doc = Document(io.BytesIO(data))
                return "\n".join(p.text for p in doc.paragraphs)
            except ImportError:
                return f"[DOCX: {filename} — python-docx not installed]"

        if ext in ("xlsx", "xls") or "spreadsheetml" in mime_type:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
                lines = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        lines.append("\t".join(str(c) if c is not None else "" for c in row))
                return "\n".join(lines)
            except ImportError:
                return f"[XLSX: {filename} — openpyxl not installed]"

        if mime_type.startswith("image/"):
            return f"[Image: {filename}]"

        return f"[Attachment: {filename}]"
    except Exception as e:
        log.warning("File extraction failed for %s: %s", filename, e)
        return f"[{filename} — extraction error]"


@router.post("/technician/assist", response_model=TechnicianAssistResponse)
async def assist_technician(
    text: str = Form(...),
    session_id: Optional[str] = Form(None),
    ticket_context: Optional[str] = Form(None),
    files: List[UploadFile] = File(default=[]),
):
    """
    Provide assistance to a technician. Accepts multipart/form-data with optional file attachments.
    """
    try:
        # Build augmented message with file content
        augmented = text
        for f in files[:3]:  # max 3 files
            data = await f.read()
            if len(data) > 10 * 1024 * 1024:
                continue  # skip oversized files silently
            extracted = _extract_file_text(f.filename or "", data, f.content_type or "")
            if extracted:
                augmented += f"\n\n[Attached: {f.filename}]\n{extracted}"

        if ticket_context:
            augmented = f"[Ticket context]\n{ticket_context}\n\n{augmented}"

        agent = get_assistant_agent()
        result = agent.assist_technician(augmented, session_id=session_id)

        if not result.get("success"):
            return TechnicianAssistResponse(success=False, message=result.get("message"))

        return TechnicianAssistResponse(
            success=True,
            session_id=result.get("session_id"),
            ticket_number=result.get("ticket_number"),
            analysis=result.get("analysis"),
            solution=result.get("solution"),
            sources=[Source(**s) for s in result.get("sources", [])],
            follow_up_questions=result.get("follow_up_questions", []),
            original_query=result.get("original_query"),
        )

    except Exception as e:
        log.error("Error in technician assistance: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


# ── Technician Availability / Attendance ──────────────────────────────────────

VALID_STATUSES = ("available", "wfh", "on_leave", "half_day", "offline", "out_of_office", "away", "busy")

STATUS_LABELS = {
    "available":    "Available",
    "wfh":          "Working from Home",
    "on_leave":     "On Leave",
    "half_day":     "Half Day",
    "offline":      "Offline",
    "out_of_office": "Out of Office",
    "away":         "Away",
    "busy":         "Busy",
}


class AvailabilityUpdateRequest(BaseModel):
    status: str
    notes: Optional[str] = None
    action: Optional[str] = None  # "punch_in" | "punch_out" | "status_change"


@router.get("/me/availability")
def get_my_availability(payload: dict = Depends(require_tech)):
    """Return current technician's status and today's attendance record."""
    tech_id = payload.get("sub")
    db = get_db_connection()

    rows = db.execute_query(
        "SELECT tech_id, tech_name, status, current_workload FROM technician_data WHERE tech_id=%s LIMIT 1",
        (tech_id,),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Technician not found")
    tech = rows[0]

    today = datetime.now(timezone.utc).date()
    att_rows = db.execute_query(
        "SELECT id, date, punch_in, punch_out, status, notes FROM technician_attendance "
        "WHERE tech_id=%s AND date=%s ORDER BY id DESC LIMIT 1",
        (tech_id, today),
    )
    attendance = att_rows[0] if att_rows else None

    return {
        "tech_id": tech["tech_id"],
        "tech_name": tech["tech_name"],
        "status": tech["status"],
        "status_label": STATUS_LABELS.get(tech["status"], tech["status"]),
        "current_workload": tech["current_workload"],
        "today": {
            "punch_in":  attendance["punch_in"].isoformat() if attendance and attendance.get("punch_in") else None,
            "punch_out": attendance["punch_out"].isoformat() if attendance and attendance.get("punch_out") else None,
            "status":    attendance["status"] if attendance else None,
            "notes":     attendance["notes"] if attendance else None,
        } if attendance else None,
    }


@router.patch("/me/availability")
def update_my_availability(req: AvailabilityUpdateRequest, payload: dict = Depends(require_tech)):
    """Update technician availability status and optionally punch in/out."""
    if req.status not in VALID_STATUSES:
        raise HTTPException(status_code=422, detail=f"Invalid status. Must be one of: {', '.join(VALID_STATUSES)}")

    tech_id = payload.get("sub")
    db = get_db_connection()

    # Update technician status
    db.execute_query(
        "UPDATE technician_data SET status=%s WHERE tech_id=%s",
        (req.status, tech_id),
        fetch=False,
    )

    now = datetime.now(timezone.utc)
    today = now.date()

    # Fetch or create today's attendance record
    att_rows = db.execute_query(
        "SELECT id, punch_in, punch_out FROM technician_attendance WHERE tech_id=%s AND date=%s ORDER BY id DESC LIMIT 1",
        (tech_id, today),
    )

    action = req.action or ("punch_in" if req.status == "available" and not att_rows else "status_change")

    if not att_rows:
        # First check-in today
        punch_in = now if req.status == "available" else None
        db.execute_query(
            "INSERT INTO technician_attendance (tech_id, date, punch_in, status, notes) VALUES (%s,%s,%s,%s,%s)",
            (tech_id, today, punch_in, req.status, req.notes),
            fetch=False,
        )
    else:
        att = att_rows[0]
        if action == "punch_in" and not att.get("punch_in"):
            db.execute_query(
                "UPDATE technician_attendance SET punch_in=%s, status=%s, notes=%s WHERE id=%s",
                (now, req.status, req.notes, att["id"]),
                fetch=False,
            )
        elif action == "punch_out":
            db.execute_query(
                "UPDATE technician_attendance SET punch_out=%s, status=%s, notes=%s WHERE id=%s",
                (now, req.status, req.notes, att["id"]),
                fetch=False,
            )
        else:
            db.execute_query(
                "UPDATE technician_attendance SET status=%s, notes=%s WHERE id=%s",
                (req.status, req.notes, att["id"]),
                fetch=False,
            )

    return {
        "success": True,
        "status": req.status,
        "status_label": STATUS_LABELS.get(req.status, req.status),
        "message": f"Status updated to {STATUS_LABELS.get(req.status, req.status)}",
    }


@router.get("/me/attendance")
def get_my_attendance(payload: dict = Depends(require_tech)):
    """Return last 14 days of attendance records for the current technician."""
    tech_id = payload.get("sub")
    db = get_db_connection()
    rows = db.execute_query(
        "SELECT date, punch_in, punch_out, status, notes FROM technician_attendance "
        "WHERE tech_id=%s AND date >= CURRENT_DATE - INTERVAL '14 days' ORDER BY date DESC",
        (tech_id,),
    )
    return {
        "attendance": [
            {
                "date":      str(r["date"]),
                "punch_in":  r["punch_in"].isoformat() if r.get("punch_in") else None,
                "punch_out": r["punch_out"].isoformat() if r.get("punch_out") else None,
                "status":    r["status"],
                "status_label": STATUS_LABELS.get(r["status"], r["status"]),
                "notes":     r.get("notes"),
            }
            for r in rows
        ]
    }
